"""Generate FuelChain Bolivia project report (Word)."""
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

OUT = Path(__file__).resolve().parent / "Informe_FuelChain_Bolivia.docx"


def set_run_font(run, *, bold=False, size=11, color=None):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), "Calibri")
    if color:
        run.font.color.rgb = color


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        set_run_font(run, bold=True, size=16 if level == 1 else 13 if level == 2 else 12)
    return h


def add_para(doc, text, *, bold=False, size=11, space_after=8):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, bold=bold, size=size)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(item, style="List Bullet")
        for run in p.runs:
            set_run_font(run, size=11)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for p in hdr[i].paragraphs:
            for run in p.runs:
                set_run_font(run, bold=True, size=10)
    for r_i, row in enumerate(rows):
        cells = table.rows[r_i + 1].cells
        for c_i, val in enumerate(row):
            cells[c_i].text = str(val)
            for p in cells[c_i].paragraphs:
                for run in p.runs:
                    set_run_font(run, size=10)
    doc.add_paragraph()


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Pt(72)
    section.bottom_margin = Pt(72)
    section.left_margin = Pt(72)
    section.right_margin = Pt(72)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("FuelChain Bolivia")
    set_run_font(r, bold=True, size=28, color=RGBColor(0x10, 0x20, 0x28))

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("Cada litro. Cada movimiento. Cada evidencia.")
    set_run_font(r, bold=True, size=14, color=RGBColor(0xD3, 0x5A, 0x00))

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run(
        f"Informe técnico del proyecto · Buildathon · {date.today().isoformat()}\n"
        "Documento DEMO / FUELCHAIN ABSTRACTION"
    )
    set_run_font(r, size=10, color=RGBColor(0x4D, 0x5F, 0x68))

    add_para(
        doc,
        "Este informe describe el sistema construido, su arquitectura, dominio de negocio, "
        "capacidades demostrables, límites éticos/técnicos y cómo operarlo en vivo ante un jurado.",
    )

    # 1
    add_heading(doc, "1. Resumen ejecutivo", 1)
    add_para(
        doc,
        "FuelChain Bolivia es una plataforma de trazabilidad, reconciliación y auditoría de la "
        "cadena logística de combustible importado. El centro del modelo es el lote (FuelBatch), "
        "identificado como FC-BO-YYYY-######. El sistema registra custodia, documentos, calidad, "
        "mediciones (hoy simuladas), discrepancias, riesgo explicable, casos de auditoría humana "
        "y evidencia digital anclada en blockchain.",
    )
    add_para(
        doc,
        "Importante: es un prototipo de hackathon. No es un sistema oficial de YPFB, ANH ni Aduana. "
        "No afirma que blockchain demuestre la existencia física de litros. Las anomalías son "
        "señales para auditoría, no sentencias de robo o corrupción. La IA explica; el auditor decide.",
    )

    # 2
    add_heading(doc, "2. Problema que resuelve", 1)
    add_para(
        doc,
        "En la importación de combustibles intervienen varios actores (proveedor, importador, "
        "transporte, aduana, depósito, laboratorio, auditor). La información suele estar fragmentada "
        "en documentos y bases aisladas. Eso dificulta responder con evidencia: qué lote es, quién "
        "lo tuvo, cuánto se declaró vs. cuánto se midió, qué documentos respaldan el movimiento y "
        "si hay discrepancia que merezca revisión humana.",
    )
    add_bullets(
        doc,
        [
            "Unificar el seguimiento en un pasaporte digital por lote.",
            "Comparar volúmenes declarados, recibidos y medidos.",
            "Generar señales de discrepancia sin culpar automáticamente.",
            "Asistir al auditor con explicación de riesgo.",
            "Anclar hashes de eventos en una capa tamper-evident (blockchain).",
        ],
    )

    # 3
    add_heading(doc, "3. Alcance y no-alcance", 1)
    add_heading(doc, "3.1 Incluido en el prototipo", 2)
    add_bullets(
        doc,
        [
            "Monorepo pnpm con web (Next.js), API (NestJS), Prisma/PostgreSQL, contratos Hardhat.",
            "Dominio FuelBatch: custodia, autorizaciones, transporte, aduana, documentos, calidad, tanques, anomalías, auditorías.",
            "Dashboard web: Resumen, Lotes, Discrepancias, Auditorías, Evidencia.",
            "Seed DEMO con tres lotes narrativos (181, 182, 184).",
            "Anclado on-chain en vivo sobre Hardhat local (botón «Anclar ahora»).",
            "Mediciones vía SIMULATOR (ESP32 físico diferido).",
        ],
    )
    add_heading(doc, "3.2 Explicitamente fuera de alcance", 2)
    add_bullets(
        doc,
        [
            "Integraciones reales con APIs gubernamentales (no inventadas).",
            "Certificación metrológica de sensores.",
            "Mainnet / fondos reales.",
            "Autenticación productiva (JWT placeholder).",
            "IPFS productivo (storage local P0).",
        ],
    )

    # 4
    add_heading(doc, "4. Arquitectura", 1)
    add_para(
        doc,
        "Separación de capas: PostgreSQL es la fuente de verdad operacional; blockchain es índice "
        "de integridad de eventos digitales; IoT aporta mediciones a reconciliar; la API orquesta; "
        "la web opera y demuestra.",
    )
    add_table(
        doc,
        ["Capa", "Tecnología", "Rol"],
        [
            ["Web", "Next.js 15 + Tailwind", "Consola operador / demo jurado"],
            ["API", "NestJS 11", "Dominio, reconciliación, anclado"],
            ["DB", "PostgreSQL + Prisma", "Estado operacional"],
            ["Chain", "Solidity + Hardhat + viem", "Evidencia tamper-evident"],
            ["IoT", "MQTT / Mosquitto / simulator", "Mediciones DEMO"],
            ["Monorepo", "pnpm + Turborepo", "Gestión de paquetes y builds"],
        ],
    )
    add_para(doc, "Estructura de carpetas principal:", bold=True)
    add_bullets(
        doc,
        [
            "apps/web — interfaz FuelChain",
            "apps/api — API NestJS",
            "prisma/ — schema, migraciones, seed",
            "contracts/ — FuelChain.sol, Hardhat, deploy",
            "packages/ — shared / config",
            "iot/ — simulador",
            "docs/ — arquitectura, proceso Bolivia, guía demo, este informe",
        ],
    )

    # 5
    add_heading(doc, "5. Modelo de dominio", 1)
    add_para(
        doc,
        "La unidad central es FuelBatch. A su alrededor cuelgan eventos de custodia, autorizaciones "
        "de importación, transportes, registros aduaneros, documentos (hash), calidad/lab, "
        "mediciones de tanque, anomalías, casos de auditoría y anclas blockchain.",
    )
    add_table(
        doc,
        ["Concepto", "Descripción"],
        [
            ["FuelBatch", "Lote FC-BO-… con producto, volumen, origen, riesgo, estado"],
            ["CustodyEvent", "Eslabón de la cadena (creación, tránsito, frontera, recepción, etc.)"],
            ["Anomaly", "Discrepancia cuantitativa/documental — señal, no veredicto"],
            ["AuditCase", "Caso humano abierto a partir de anomalías"],
            ["BlockchainAnchor", "Índice off-chain de txHash / dataHash / bloque"],
            ["Measurement", "Lectura SIMULATOR (o IoT futuro) asociada a tanque/lote"],
        ],
    )

    # 6
    add_heading(doc, "6. Datos DEMO (seed)", 1)
    add_table(
        doc,
        ["Código", "Riesgo / estado", "Historia"],
        [
            ["FC-BO-2026-000181", "LOW / COMPLETED", "Camino feliz"],
            ["FC-BO-2026-000182", "MEDIUM / IN_TRANSIT", "En tránsito"],
            [
                "FC-BO-2026-000184",
                "HIGH / AUDIT_REQUIRED",
                "Caso estrella: Declared 100000 → Received 99900 → Stored 98700 → SIMULATOR 98650 + anomalía + auditoría",
            ],
        ],
    )
    add_para(
        doc,
        "Frase clave ante el jurado: no decimos «robo». Decimos ANOMALY / DISCREPANCY: puede ser "
        "medición, calibración, temperatura, documentación u otros factores.",
    )

    # 7
    add_heading(doc, "7. API (principales endpoints)", 1)
    add_table(
        doc,
        ["Método / ruta", "Uso"],
        [
            ["GET /health", "Salud API + DB"],
            ["GET /dashboard/kpis", "KPIs del resumen"],
            ["GET /batches", "Listado y filtros"],
            ["GET /batches/:idOrCode/passport", "Pasaporte digital del lote"],
            ["GET /anomalies", "Centro de discrepancias"],
            ["GET /audits", "Casos de auditoría"],
            ["GET /blockchain/anchors", "Índice de evidencia"],
            ["GET /blockchain/status", "Estado live de Hardhat/contrato"],
            ["POST /blockchain/anchor", "Anclar evento on-chain en vivo"],
            ["GET /blockchain/verify/:txHash", "Verificar receipt"],
        ],
    )

    # 8
    add_heading(doc, "8. Interfaz web", 1)
    add_para(
        doc,
        "Dirección visual: «despacho diurno / patio de tanques». Tipografía Chivo + IBM Plex Sans. "
        "Marca FuelChain BOLIVIA dominante en el mástil. Páginas: Resumen (medidor de tanque), "
        "Lotes, Discrepancias, Auditorías, Evidencia (panel Anclar ahora).",
    )
    add_bullets(
        doc,
        [
            "http://localhost:3000/ — Resumen",
            "http://localhost:3000/batches — Lotes",
            "http://localhost:3000/batches/FC-BO-2026-000184 — Pasaporte del caso estrella",
            "http://localhost:3000/anomalies — Discrepancias",
            "http://localhost:3000/audits — Auditorías",
            "http://localhost:3000/blockchain — Evidencia + anclado en vivo",
        ],
    )

    # 9
    add_heading(doc, "9. Blockchain (demo en vivo)", 1)
    add_para(
        doc,
        "Contrato FuelChain.sol: función anchorEvidence(batchId, eventKind, dataHash). Solo almacena "
        "identificadores/hashes — nunca PDFs ni PII. La API usa viem contra Hardhat (chainId 31337). "
        "Tras confirmar la transacción, persiste BlockchainAnchor en Postgres con txHash y bloque.",
    )
    add_para(doc, "Arranque en vivo:", bold=True)
    add_bullets(
        doc,
        [
            "Terminal A: pnpm contracts:node  (dejar corriendo)",
            "Terminal B: pnpm contracts:compile && pnpm --filter @fuelchain/contracts run deploy",
            "Copiar FUELCHAIN_CONTRACT_ADDRESS al .env (o dejar deployments/localhost.json)",
            "BLOCKCHAIN_PRIVATE_KEY = cuenta #0 de Hardhat (solo local DEMO)",
            "Reiniciar API → GET /blockchain/status debe devolver live: true",
            "UI Evidencia → Anclar ahora → mostrar txHash real",
        ],
    )
    add_para(
        doc,
        "Respuesta canónica: blockchain no prueba que el litro exista. Prueba que este evento y este "
        "hash quedaron registrados de forma resistente a modificación.",
    )

    # 10
    add_heading(doc, "10. Cómo arrancar el proyecto", 1)
    add_para(doc, "Requisitos: Node.js ≥ 20, pnpm (nunca npm/npx), Docker para Postgres.", bold=False)
    add_bullets(
        doc,
        [
            "docker compose up -d postgres",
            "pnpm install",
            "pnpm db:generate && pnpm db:migrate && pnpm db:seed",
            "pnpm --filter @fuelchain/api dev   → :3001",
            "pnpm --filter @fuelchain/web dev   → :3000",
            "Opcional blockchain: pnpm contracts:node + deploy (ver sección 9)",
        ],
    )
    add_para(
        doc,
        "Variables críticas en .env (nunca commitear secretos reales): DATABASE_URL, "
        "NEXT_PUBLIC_API_URL, CHAIN_RPC_URL, FUELCHAIN_CONTRACT_ADDRESS, BLOCKCHAIN_PRIVATE_KEY. "
        "Plantilla: .env.example.",
    )

    # 11
    add_heading(doc, "11. Guion de demo (5–8 min)", 1)
    add_bullets(
        doc,
        [
            "Resumen: volumen en custodia + medidor tanque + KPIs.",
            "Lotes: contrastar 181 (feliz), 182 (tránsito), 184 (auditoría).",
            "Pasaporte 184: custodia + diferencia de volúmenes + discrepancias.",
            "Auditorías: explicación asistida (mock); decisión humana.",
            "Evidencia: Anclar ahora → txHash + bloque en vivo.",
        ],
    )

    # 12
    add_heading(doc, "12. Principios éticos y de comunicación", 1)
    add_table(
        doc,
        ["Evitar", "Preferir"],
        [
            ["«Robo / corrupción detectados»", "Discrepancia / anomalía para auditoría"],
            ["«Integramos YPFB/ANH vía API oficial»", "Abstracción; sin inventar endpoints"],
            ["«El sensor prueba los litros»", "Medición a reconciliar (SIMULATOR hoy)"],
            ["«La IA decide el fraude»", "La IA explica; el auditor decide"],
            ["«Blockchain evita el robo»", "Capa de evidencia compartida tamper-evident"],
        ],
    )

    # 13
    add_heading(doc, "13. Estado de fases y trabajo futuro", 1)
    add_table(
        doc,
        ["Área", "Estado"],
        [
            ["Arquitectura / monorepo", "Hecho"],
            ["Prisma + Postgres + seed", "Hecho"],
            ["API dominio + dashboard", "Hecho"],
            ["UI web + design system", "Hecho"],
            ["Blockchain live Hardhat", "Hecho (local DEMO)"],
            ["ESP32 físico", "Diferido"],
            ["Auth / roles", "Pendiente P1"],
            ["Testnet pública (Base Sepolia, etc.)", "Opcional"],
            ["IPFS documentos", "P2"],
            ["IA proveedor real", "Opcional (hoy mock)"],
        ],
    )

    # 14
    add_heading(doc, "14. Repositorio y licencia de uso DEMO", 1)
    add_para(
        doc,
        "Código destinado a demostración en Buildathon. Etiquetar siempre DEMO / ASSUMPTION / "
        "FUELCHAIN ABSTRACTION donde corresponda. No desplegar la private key de Hardhat fuera de "
        "entorno local. No usar claves con fondos reales en el prototipo.",
    )
    add_para(
        doc,
        "Documentación complementaria en el repo: docs/architecture.md, docs/bolivia-fuel-process.md, "
        "docs/demo.md, apps/web/DESIGN_SYSTEM.md.",
    )

    add_heading(doc, "15. Conclusión", 1)
    add_para(
        doc,
        "FuelChain Bolivia demuestra un flujo completo de integridad de cadena: lote → custodia → "
        "medición → discrepancia → riesgo → auditoría humana → evidencia on-chain. El valor no está "
        "en afirmar certeza física absoluta, sino en hacer visible, reconciliable y anclable cada "
        "movimiento relevante del combustible importado.",
    )

    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("— Fin del informe —\nFuelChain Bolivia · Buildathon")
    set_run_font(r, size=10, color=RGBColor(0x4D, 0x5F, 0x68))

    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
