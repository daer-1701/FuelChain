"""Generate Word: credenciales DEMO + lista de compra sensores cisterna/tanque."""
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

OUT = Path(__file__).resolve().parent / "FuelChain_Credenciales_y_Hardware.docx"


def set_run_font(run, *, bold=False, size=11, color=None):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), "Calibri")
    if color:
        run.font.color.rgb = color


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        set_run_font(
            run,
            bold=True,
            size=16 if level == 1 else 13 if level == 2 else 12,
        )
    return h


def add_para(doc, text, *, bold=False, size=11, space_after=8):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, bold=bold, size=size)
    p.paragraph_format.space_after = Pt(space_after)
    return p


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for run in p.runs:
                set_run_font(run, bold=True, size=10)
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            cell.text = str(val)
            for p in cell.paragraphs:
                for run in p.runs:
                    set_run_font(run, size=10)
    doc.add_paragraph()


def main():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Pt(64)
        s.bottom_margin = Pt(64)
        s.left_margin = Pt(64)
        s.right_margin = Pt(64)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("FuelChain Bolivia")
    set_run_font(r, bold=True, size=26, color=RGBColor(0x10, 0x20, 0x28))

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("Credenciales DEMO y componentes para cisternas / tanques")
    set_run_font(r, bold=True, size=14, color=RGBColor(0xD3, 0x5A, 0x00))

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run(
        f"Piloto Cochabamba · {date.today().isoformat()} · Solo entorno DEMO — no producción"
    )
    set_run_font(r, size=10, color=RGBColor(0x4D, 0x5F, 0x68))

    add_para(
        doc,
        "Este documento resume (1) usuarios de acceso al sistema demo y "
        "(2) qué comprar/instalar para medir nivel y proxies de calidad en "
        "cisternas y tanques de estación, respetando atmósfera explosiva "
        "(vapores de gasolina/diésel).",
    )

    # ── Credenciales ─────────────────────────────────────────────────────────
    add_heading(doc, "1. Credenciales de acceso (DEMO)", 1)
    add_para(
        doc,
        "URL web: http://localhost:3000/login  ·  API: http://localhost:3001",
        bold=True,
    )
    add_para(
        doc,
        "Contraseña única para todos los usuarios DEMO:  demo123",
        bold=True,
    )
    add_para(
        doc,
        "El mapa público /mapa no requiere login. El botón "
        "«Ver surtidores (sin login)» en la pantalla de acceso abre el mapa.",
    )

    add_table(
        doc,
        ["Correo", "Nombre", "Rol", "Para qué sirve"],
        [
            [
                "chofer@fuelchain.bo",
                "Chofer cisterna DEMO",
                "TRANSPORTER",
                "Emitir bastón QR / custodia en ruta",
            ],
            [
                "estacion@fuelchain.bo",
                "Encargado EESS Cala Cala",
                "STATION_STAFF",
                "Aceptar entrega, actualizar tanque",
            ],
            [
                "deposito@fuelchain.bo",
                "Operador Depósito DEMO",
                "DEPOT_OPERATOR",
                "Carga / recepción en depósito",
            ],
            [
                "importador@fuelchain.bo",
                "Importador DEMO CBBA",
                "IMPORTER",
                "Lotes y autorizaciones",
            ],
            [
                "lab@fuelchain.bo",
                "Laboratorio DEMO",
                "LAB",
                "Certificados / muestreo (demo)",
            ],
            [
                "auditor@fuelchain.bo",
                "Auditor DEMO",
                "AUDITOR",
                "Discrepancias y casos de auditoría",
            ],
            [
                "auditor.demo@fuelchain.bo",
                "Auditor DEMO (alias)",
                "AUDITOR",
                "Misma función (alias legacy)",
            ],
            [
                "anh@fuelchain.bo",
                "Verificador ANH (read-only DEMO)",
                "VERIFIER",
                "Consulta pasaporte / verify",
            ],
            [
                "ciudadano@fuelchain.bo",
                "Ciudadano mapa CBBA",
                "VERIFIER",
                "Consulta (opcional; mapa es público)",
            ],
        ],
    )

    add_para(
        doc,
        "Advertencia: credenciales solo para hackathon/demo local. "
        "No usar en internet ni con datos reales. Hash DEMO = scrypt local "
        "(no IdP productivo).",
        size=10,
    )

    # ── Hardware ─────────────────────────────────────────────────────────────
    add_heading(doc, "2. Qué comprar para cisternas y tanques", 1)
    add_para(
        doc,
        "Regla de oro: dentro del vapor del tanque/cisterna SOLO entra una "
        "sonda certificada Ex ia / ATEX / IECEx. El cerebro (gateway, ESP32, "
        "módem 4G, baterías) vive en zona segura (cabina, cuarto eléctrico, "
        "oficina de playa).",
        bold=True,
    )

    add_heading(doc, "2.1 Arquitectura de compra", 2)
    add_para(
        doc,
        "Zona peligrosa → sonda Ex ia → cable IS → barrera intrínseca "
        "(zona segura) → gateway FuelChain → API / cola offline.",
    )

    add_heading(doc, "2.2 Lista de compra — TANQUE de estación (EESS)", 2)
    add_table(
        doc,
        ["Componente", "Especificación a pedir", "Cant. tip.", "Notas"],
        [
            [
                "Sonda de nivel",
                "Magnetostrictiva o radar/ultrasónico Ex ia, Zone 0/1, apta gasolina/diésel; preferible con interfase agua",
                "1 por tanque",
                "Estándar EESS; no comprar “maker”",
            ],
            [
                "Barrera / aislador IS",
                "Compatible con el lazo del sensor (Zener o activo), marcado Ex",
                "1 por lazo",
                "Se instala en gabinete zona segura",
            ],
            [
                "Sensor temperatura",
                "RTD/termistor Ex ia o integrado en la sonda",
                "1",
                "Proxy de calidad",
            ],
            [
                "Detección de agua",
                "Interfase agua-combustible (magnetostrictivo o flotador certificado)",
                "1",
                "Crítico en subterráneo",
            ],
            [
                "Densidad / dieléctrico (opcional)",
                "Sonda combinada Ex ia si el fabricante la ofrece",
                "0–1",
                "Proxy; no reemplaza lab",
            ],
            [
                "Gateway / edge",
                "PLC industrial o PC industrial / ESP32 solo en zona segura + 4G/WiFi estación",
                "1 por sitio",
                "Store-and-forward si no hay señal",
            ],
            [
                "Gabinete + UPS",
                "IP65, puesta a tierra, alimentación estabilizada",
                "1",
                "Fuera de zona 0/1",
            ],
            [
                "Cableado IS",
                "Cable intrínsecamente seguro, segregado, bonding",
                "Según plano",
                "Instalador autorizado",
            ],
        ],
    )

    add_heading(doc, "2.3 Lista de compra — CISTERNA", 2)
    add_table(
        doc,
        ["Componente", "Especificación a pedir", "Cant. tip.", "Notas"],
        [
            [
                "Sonda de nivel cisterna",
                "Ex ia para compartimento de producto (Zone 0/1)",
                "1 por compartimento",
                "Nunca electrónica WiFi dentro del vapor",
            ],
            [
                "Barrera IS",
                "En gabinete de cabina / chasis zona segura",
                "1 por lazo",
                "Obligatoria",
            ],
            [
                "Logger / gateway cabina",
                "Equipo industrial o gateway 4G/NB-IoT con SD; GPS en cabina",
                "1",
                "Buffer offline (sin señal en ruta Arica–Bolivia)",
            ],
            [
                "GPS + antena",
                "En cabina (fuera de vapor)",
                "1",
                "Trazabilidad de ruta",
            ],
            [
                "Teléfono endurecido / PWA",
                "Android con cámara para QR de custodia",
                "1 por chofer",
                "Cola offline FuelChain",
            ],
            [
                "Temperatura / agua proxy",
                "Integrados a la sonda Ex o lazos adicionales IS",
                "Según diseño",
                "Calidad proxy",
            ],
            [
                "Alimentación vehículo",
                "Convertidor 12/24 V → 24 Vdc aislado + fusibles",
                "1",
                "Zona segura",
            ],
        ],
    )

    add_heading(doc, "2.4 Familias de referencia (no endorsement)", 2)
    add_para(
        doc,
        "Al cotizar, pedir certificado ATEX/IECEx Ex ia vigente y compatibilidad "
        "con hidrocarburos líquidos. Ejemplos de familias citadas en la industria "
        "(verificar modelo exacto con el proveedor local en Bolivia/Chile):",
    )
    add_para(
        doc,
        "• Magnetostrictivos ATEX Zone 0 para EESS (nivel + agua).\n"
        "• Ultrasónicos / radar IS (p. ej. familias tipo TEK Exi, APG MNU IS u equivalentes regionales).\n"
        "• Barreras Pepperl+Fuchs / MTL / Phoenix Contact (o equivalente certificado).\n"
        "• Gateway: PLC Siemens/Schneider edge, o PC industrial; ESP32 solo demo en zona segura.",
    )

    add_heading(doc, "2.5 Presupuesto orientativo por sitio (DEMO / piloto)", 2)
    add_para(
        doc,
        "Los precios reales dependen de importación, certificación y mano de obra "
        "calificada. Órdenes de magnitud típicas de mercado industrial (USD, "
        "referenciales, no cotización):",
    )
    add_table(
        doc,
        ["Paquete", "Alcance", "Rango referencial USD"],
        [
            [
                "EESS 1 tanque",
                "Sonda Ex + barrera + gateway básico + instalación",
                "3.000 – 12.000+",
            ],
            [
                "Cisterna 1 compartimento",
                "Sonda Ex + barrera + logger/GPS cabina",
                "2.500 – 10.000+",
            ],
            [
                "Solo DEMO hackathon",
                "Sin hardware: SIMULATOR + celular QR",
                "0 (software)",
            ],
        ],
    )

    add_heading(doc, "3. Checklist de instalación segura", 1)
    for item in [
        "Clasificar zonas 0 / 1 / 2 con persona competente.",
        "Comprar solo sensores con marcado Ex ia (o equivalente) para gasolina/diésel.",
        "Instalar barrera IS en área segura; cableado segregado y bonding.",
        "Gateway / módem / baterías FUERA del compartimento de vapor.",
        "Documentar modelo, certificado, plano y pruebas.",
        "En demo actual del monorepo: usar SIMULATOR; hardware físico = piloto real.",
        "Celular del operador: leer QR en zona segura (oficina de playa), no sobre boca de carga activa.",
    ]:
        p = doc.add_paragraph(item, style="List Number")
        for run in p.runs:
            set_run_font(run, size=11)

    add_heading(doc, "4. Relación con la importación (Arica)", 1)
    add_para(
        doc,
        "La gasolina/diésel suelen descargarse de buques en Terminal Sica Sica "
        "(Arica, Chile) y viajar en cisterna a Bolivia. Por eso el logger offline "
        "en cabina + QR entre actores es crítico en tramos sin señal. Ver "
        "docs/importacion-chile-arica.md en el repositorio.",
    )

    add_heading(doc, "5. Frase para el jurado", 1)
    add_para(
        doc,
        "“Medimos cantidad y proxies de calidad en cisterna y tanque, pero la "
        "electrónica inteligente vive en zona segura. En el tanque solo entra "
        "una sonda intrínsecamente segura con barrera — porque el vapor de "
        "gasolina puede inflamarse.”",
        bold=True,
    )

    foot = doc.add_paragraph()
    r = foot.add_run(
        "FuelChain Bolivia · Documento DEMO · No es lista de compra oficial YPFB/ANH · "
        "No sustituye ingeniería de seguridad certificada."
    )
    set_run_font(r, size=9, color=RGBColor(0x4D, 0x5F, 0x68))

    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
